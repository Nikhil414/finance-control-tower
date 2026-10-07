"""Build the Excel human-review workbook for a run.

Excel is the review interface because that is where finance reviewers actually
work. The design constraint is that it must never become a second source of
truth: the workbook is generated from Postgres, reviewers fill in decision
columns, and a validated CSV round-trip carries those decisions back.

No VBA and no direct database writes from Excel. A macro-enabled workbook with
a live connection would put an uncontrolled write path into the middle of a
financial control system, and would be blocked by most corporate policy anyway.
The round-trip is slower and completely auditable:

    Postgres -> workbook -> reviewer -> review_decisions.csv -> validation -> Postgres

Data validation dropdowns constrain decisions to the contract's vocabulary, so a
typo becomes impossible rather than merely detectable.

Run: python src/build_review_workbook.py
"""

from __future__ import annotations

import json
import logging
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from contract import CONTRACT
from db import connect

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_DIR = BASE_DIR / "excel"

logger = logging.getLogger("capstone.workbook")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
LOCKED_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(bold=True, size=14)
THIN_BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)

SEVERITY_FILLS = {
    "CRITICAL": PatternFill("solid", fgColor="FFC7CE"),
    "HIGH": PatternFill("solid", fgColor="FFD9A0"),
    "MEDIUM": PatternFill("solid", fgColor="FFF2CC"),
    "LOW": PatternFill("solid", fgColor="E2EFDA"),
}

# Columns the reviewer fills in. Everything else is generated and read-only.
DECISION_COLUMNS = (
    ("decision", 26),
    ("reviewer", 18),
    ("reviewed_at", 14),
    ("reviewer_notes", 40),
    ("approved_action", 34),
    ("confirmed_loss", 15),
    ("recovered_amount", 16),
    ("false_positive_reason", 34),
    ("resolution_date", 14),
)

QUEUE_COLUMNS = (
    ("finding_id", 20),
    ("priority_score", 11),
    ("severity", 10),
    ("source_module", 15),
    ("rule_code", 30),
    ("entity_type", 12),
    ("entity_id", 20),
    ("order_id", 16),
    ("customer_id", 14),
    ("business_date", 13),
    ("risk_amount", 15),
    ("due_date", 12),
    ("approval_required", 17),
    ("status", 10),
    ("recommended_action", 50),
    ("ai_fp_likelihood", 16),
)


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.row_dimensions[row].height = 30


def _set_widths(sheet, columns) -> None:
    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def load_fp_likelihood(output_dir: Path) -> dict[str, float]:
    """Jev's false-positive triage hint, keyed by finding_id.

    Read from disk rather than the database because it is advisory metadata
    from an optional AI step, not a database-backed fact -- ai_commentary.json
    may not exist (Jev unconfigured) or may predate the current run.
    """
    path = output_dir / "ai_commentary.json"
    if not path.exists():
        return {}
    commentary = json.loads(path.read_text(encoding="utf-8"))
    return commentary.get("false_positive_likelihood", {})


def fetch_open_findings(conn, run_id: str) -> list[dict]:
    with conn.cursor() as cur:
        cur.execute("""
            SELECT finding_id, priority_score, severity, source_module, rule_code,
                   entity_type, entity_id, order_id, customer_id, business_date,
                   risk_amount, due_date, approval_required, status,
                   recommended_action, evidence_json
            FROM capstone.fact_findings
            WHERE last_seen_run_id = %s AND status NOT IN ('RESOLVED', 'FALSE_POSITIVE')
            ORDER BY priority_score DESC, risk_amount DESC
        """, (run_id,))
        columns = [description.name for description in cur.description]
        return [dict(zip(columns, row)) for row in cur.fetchall()]


def fetch_reconciliation(conn, run_id: str) -> list[dict]:
    with conn.cursor() as cur:
        cur.execute("""
            SELECT match_status, bank_txn_id, ledger_txn_id, bank_amount,
                   ledger_amount, amount_difference, bank_value_date,
                   ledger_posting_date, day_difference, aging_bucket
            FROM capstone.fact_reconciliation_results
            WHERE run_id = %s AND match_status <> 'MATCHED'
            ORDER BY ABS(COALESCE(bank_amount, ledger_amount, 0)) DESC
        """, (run_id,))
        columns = [description.name for description in cur.description]
        return [dict(zip(columns, row)) for row in cur.fetchall()]


def fetch_control_totals(conn, run_id: str) -> list[tuple]:
    """Figures a reviewer can tie the workbook back to Postgres with.

    Without these the workbook is an unverifiable extract: someone comparing it
    to the dashboard would have no way to tell whether they match.
    """
    with conn.cursor() as cur:
        cur.execute("""
            SELECT
                (SELECT COUNT(*) FROM capstone.fact_findings WHERE last_seen_run_id = %(run)s),
                (SELECT COALESCE(SUM(risk_amount), 0) FROM capstone.fact_findings WHERE last_seen_run_id = %(run)s),
                (SELECT deduplicated_exposure FROM capstone.vw_run_exposure WHERE run_id = %(run)s),
                (SELECT COUNT(*) FROM capstone.fact_findings WHERE last_seen_run_id = %(run)s AND severity = 'CRITICAL'),
                (SELECT COUNT(*) FROM capstone.fact_findings WHERE last_seen_run_id = %(run)s AND approval_required),
                (SELECT match_rate_pct FROM capstone.vw_reconciliation_summary),
                (SELECT breaks FROM capstone.vw_reconciliation_summary),
                (SELECT data_quality_score_pct FROM capstone.vw_data_quality_score),
                (SELECT COUNT(*) FROM staging.orders),
                (SELECT COUNT(*) FROM staging.payments)
        """, {"run": run_id})
        row = cur.fetchone()

    return [
        ("Findings in this run", row[0], "capstone.fact_findings"),
        ("Gross exposure (INR)", float(row[1]), "SUM(risk_amount) - double-counts shared orders"),
        ("Deduplicated exposure (INR)", float(row[2] or 0), "capstone.vw_run_exposure - the defensible figure"),
        ("Critical findings", row[3], "severity = 'CRITICAL'"),
        ("Findings needing approval", row[4], "approval_required = true"),
        ("Reconciliation match rate %", float(row[5] or 0), "capstone.vw_reconciliation_summary"),
        ("Reconciliation breaks", row[6], "capstone.vw_reconciliation_summary"),
        ("Data quality score %", float(row[7] or 0), "capstone.vw_data_quality_score"),
        ("Source orders loaded", row[8], "staging.orders"),
        ("Source payments loaded", row[9], "staging.payments"),
    ]


# ---------------------------------------------------------------- sheets


def _sheet_instructions(workbook, run_id: str, finding_count: int) -> None:
    sheet = workbook.create_sheet("Instructions")
    sheet.column_dimensions["A"].width = 110

    content = [
        ("Finance Review Workbook", TITLE_FONT),
        ("", None),
        (f"Run ID: {run_id}", Font(bold=True)),
        (f"Generated: {datetime.now(timezone.utc).strftime('%d %B %Y %H:%M UTC')}", None),
        (f"Open cases awaiting review: {finding_count}", None),
        ("", None),
        ("How to use this workbook", Font(bold=True, size=12)),
        ("", None),
        ("1. Work the 'Review Queue' sheet from the top. It is sorted by priority, which blends", None),
        ("   financial exposure, severity, case age and customer segment.", None),
        ("2. Open 'Finding Evidence' and locate your finding_id. The evidence there is sufficient", None),
        ("   to reproduce the flagged amount by hand. If you cannot reproduce it, the correct", None),
        ("   decision is MORE_INFORMATION_REQUIRED, not an approval.", None),
        ("3. Fill in ONLY the amber columns. Grey columns are generated and are overwritten on", None),
        ("   the next run.", None),
        ("4. Save as CSV to excel/review_decisions.csv, then run: python src/import_reviews.py", None),
        ("", None),
        ("What each decision means", Font(bold=True, size=12)),
        ("", None),
        ("FALSE_POSITIVE              The rule fired but there is no real exception.", None),
        ("                            Requires false_positive_reason. Cannot carry a confirmed loss.", None),
        ("VALID_EXCEPTION             A real exception. Requires reviewer_notes and confirmed_loss", None),
        ("                            (enter 0 if the amount turned out to be recoverable in full).", None),
        ("MORE_INFORMATION_REQUIRED   Cannot be concluded yet. Returns the case to review.", None),
        ("ACTION_APPROVED             Remediation authorised. Requires approved_action.", None),
        ("RESOLVED                    Case closed. Requires resolution_date.", None),
        ("", None),
        ("Estimated exposure is not confirmed loss", Font(bold=True, size=12)),
        ("", None),
        ("risk_amount is what a deterministic rule estimated might be at stake. It is NOT a loss.", None),
        ("confirmed_loss is what you verify actually happened, and the two are reported separately", None),
        ("and never added together. Entering risk_amount into confirmed_loss without verifying it", None),
        ("defeats the entire purpose of this review step.", None),
        ("", None),
        ("recovered_amount may not exceed confirmed_loss. The importer rejects the file if it does.", None),
        ("", None),
        ("Rules the importer enforces", Font(bold=True, size=12)),
        ("", None),
        ("- Decisions must be one of the values above, spelled exactly. Use the dropdown.", None),
        ("- finding_id must exist in this run.", None),
        ("- The status change implied by your decision must be a legal transition.", None),
        ("- Required fields for your chosen decision must be filled in.", None),
        ("- An invalid file is rejected in full. Nothing is imported partially, because a", None),
        ("  half-applied set of decisions is harder to unpick than a rejected file.", None),
        ("", None),
        ("Decisions are append-only once imported. To correct one, submit a new decision for the", None),
        ("same finding: the later review supersedes the earlier one, and both are kept.", None),
    ]

    for index, (text, font) in enumerate(content, start=1):
        cell = sheet.cell(row=index, column=1, value=text)
        if font:
            cell.font = font


def _sheet_review_queue(workbook, findings: list[dict]) -> None:
    sheet = workbook.create_sheet("Review Queue")
    headers = [name for name, _ in QUEUE_COLUMNS] + [name for name, _ in DECISION_COLUMNS]
    sheet.append(headers)
    _set_widths(sheet, QUEUE_COLUMNS + DECISION_COLUMNS)
    _style_header(sheet)

    for finding in findings:
        sheet.append(
            [finding[name] for name, _ in QUEUE_COLUMNS]
            + [None] * len(DECISION_COLUMNS)
        )

    generated_count = len(QUEUE_COLUMNS)
    for row in range(2, len(findings) + 2):
        severity = sheet.cell(row=row, column=3).value
        for column in range(1, generated_count + 1):
            cell = sheet.cell(row=row, column=column)
            cell.fill = SEVERITY_FILLS.get(severity, LOCKED_FILL) if column == 3 else LOCKED_FILL
            cell.border = THIN_BORDER
        for column in range(generated_count + 1, generated_count + len(DECISION_COLUMNS) + 1):
            cell = sheet.cell(row=row, column=column)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER

        sheet.cell(row=row, column=11).number_format = '#,##0.00'
        sheet.cell(row=row, column=generated_count).number_format = '0.0%'
        sheet.cell(row=row, column=generated_count + 6).number_format = '#,##0.00'
        sheet.cell(row=row, column=generated_count + 7).number_format = '#,##0.00'

    # Constrain decisions to the contract's vocabulary. A dropdown makes a typo
    # impossible rather than merely detectable at import time.
    decisions = ",".join(CONTRACT["decision_requirements"].keys() - {"note"})
    validation = DataValidation(
        type="list", formula1=f'"{decisions}"', allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid decision",
        error="Pick a decision from the list. Free text is rejected on import.",
    )
    sheet.add_data_validation(validation)
    decision_letter = get_column_letter(generated_count + 1)
    validation.add(f"{decision_letter}2:{decision_letter}{max(len(findings) + 1, 2)}")

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions


def _sheet_finding_evidence(workbook, findings: list[dict]) -> None:
    """One row per evidence key, so a reviewer can read the calculation.

    Long form rather than one JSON blob per finding: a reviewer needs to compare
    two operands and a difference, and that is unreadable inside a JSON string
    in a single cell.
    """
    sheet = workbook.create_sheet("Finding Evidence")
    columns = (("finding_id", 20), ("rule_code", 30), ("risk_amount", 15),
               ("evidence_key", 30), ("evidence_value", 44))
    sheet.append([name for name, _ in columns])
    _set_widths(sheet, columns)
    _style_header(sheet)

    for finding in findings:
        evidence = finding.get("evidence_json") or {}
        for key, value in sorted(evidence.items()):
            sheet.append([
                finding["finding_id"], finding["rule_code"],
                float(finding["risk_amount"]), key, str(value),
            ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _sheet_reconciliation(workbook, breaks: list[dict]) -> None:
    sheet = workbook.create_sheet("Reconciliation")
    columns = (("match_status", 22), ("bank_txn_id", 16), ("ledger_txn_id", 16),
               ("bank_amount", 15), ("ledger_amount", 15), ("amount_difference", 17),
               ("bank_value_date", 16), ("ledger_posting_date", 18),
               ("day_difference", 14), ("aging_bucket", 12))
    sheet.append([name for name, _ in columns])
    _set_widths(sheet, columns)
    _style_header(sheet)

    for row in breaks:
        sheet.append([row[name] for name, _ in columns])

    for row in range(2, len(breaks) + 2):
        for column in (4, 5, 6):
            sheet.cell(row=row, column=column).number_format = '#,##0.00'

    sheet.freeze_panes = "A2"
    if breaks:
        sheet.auto_filter.ref = sheet.dimensions


def _sheet_control_totals(workbook, totals: list[tuple], run_id: str) -> None:
    sheet = workbook.create_sheet("Control Totals")
    columns = (("measure", 34), ("value", 20), ("source", 62))
    sheet.append([name for name, _ in columns])
    _set_widths(sheet, columns)
    _style_header(sheet)

    for measure, value, source in totals:
        sheet.append([measure, value, source])

    for row in range(2, len(totals) + 2):
        value_cell = sheet.cell(row=row, column=2)
        if isinstance(value_cell.value, float):
            value_cell.number_format = '#,##0.00'

    note_row = len(totals) + 3
    sheet.cell(row=note_row, column=1, value="Why this sheet exists").font = Font(bold=True)
    for offset, line in enumerate([
        "Every figure here is queryable from PostgreSQL, named in the source column.",
        f"They must match the Power BI Control Tower for run {run_id} exactly.",
        "If a dashboard tile disagrees with this sheet, the dashboard is wrong -- the",
        "database is the source of truth, not the report.",
        "",
        "Gross and deduplicated exposure differ on purpose. One order can trip several",
        "rules, so the gross figure counts the same money more than once. Headline",
        "reporting uses the deduplicated figure.",
    ], start=1):
        sheet.cell(row=note_row + offset, column=1, value=line)


def _sheet_management_summary(workbook, run_id: str) -> None:
    """Embed the executive brief so reviewers see the same narrative as management."""
    sheet = workbook.create_sheet("Management Summary")
    sheet.column_dimensions["A"].width = 110

    brief_path = BASE_DIR / "outputs" / "executive_brief.md"
    if not brief_path.exists():
        sheet.cell(row=1, column=1,
                   value="No executive brief generated for this run.").font = Font(italic=True)
        return

    for index, line in enumerate(brief_path.read_text(encoding="utf-8").splitlines(), start=1):
        cell = sheet.cell(row=index, column=1, value=line)
        if line.startswith("# "):
            cell.font = TITLE_FONT
        elif line.startswith("## "):
            cell.font = Font(bold=True, size=12)


def _sheet_lists_and_rules(workbook) -> None:
    """The contract's vocabulary and rule registry, so a reviewer can see what a
    rule means and what exposure it claims to measure without leaving Excel."""
    sheet = workbook.create_sheet("Lists and Rules")

    sheet.cell(row=1, column=1, value="Valid decisions").font = Font(bold=True)
    sheet.cell(row=1, column=2, value="Required fields").font = Font(bold=True)
    row = 2
    for decision, required in CONTRACT["decision_requirements"].items():
        if decision == "note":
            continue
        sheet.cell(row=row, column=1, value=decision)
        sheet.cell(row=row, column=2, value=", ".join(required))
        row += 1

    row += 2
    sheet.cell(row=row, column=1, value="Severity levels").font = Font(bold=True)
    row += 1
    for level in CONTRACT["severity_levels"]:
        sheet.cell(row=row, column=1, value=level)
        sheet.cell(row=row, column=2,
                   value=f"SLA {CONTRACT['sla_days_by_severity'][level]} day(s)")
        row += 1

    row += 2
    sheet.cell(row=row, column=1, value="Rule registry").font = Font(bold=True)
    row += 1
    for header_index, header in enumerate(
        ("rule_code", "module", "default severity", "approval", "what the exposure means"),
        start=1,
    ):
        cell = sheet.cell(row=row, column=header_index, value=header)
        cell.font = Font(bold=True)
    row += 1

    for rule in CONTRACT["rules"]:
        sheet.cell(row=row, column=1, value=rule["rule_code"])
        sheet.cell(row=row, column=2, value=rule["source_module"])
        sheet.cell(row=row, column=3, value=rule["default_severity"])
        sheet.cell(row=row, column=4,
                   value="required" if rule["approval_required"] else "not required")
        sheet.cell(row=row, column=5, value=rule["exposure_basis"])
        row += 1

    for column, width in (("A", 34), ("B", 24), ("C", 18), ("D", 14), ("E", 90)):
        sheet.column_dimensions[column].width = width


def _sheet_decision_template(workbook, findings: list[dict]) -> None:
    """The exact shape import_reviews.py expects, pre-filled with finding_ids.

    A separate sheet from the queue because the importer reads a flat CSV: asking
    someone to hand-build that layout is how column-order mistakes happen.
    """
    sheet = workbook.create_sheet("Reviewer Decisions")
    columns = (("finding_id", 20),) + DECISION_COLUMNS
    sheet.append([name for name, _ in columns])
    _set_widths(sheet, columns)
    _style_header(sheet)

    for finding in findings:
        sheet.append([finding["finding_id"]] + [None] * len(DECISION_COLUMNS))

    for row in range(2, len(findings) + 2):
        sheet.cell(row=row, column=1).fill = LOCKED_FILL
        for column in range(2, len(columns) + 1):
            sheet.cell(row=row, column=column).fill = INPUT_FILL

    decisions = ",".join(CONTRACT["decision_requirements"].keys() - {"note"})
    validation = DataValidation(
        type="list", formula1=f'"{decisions}"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid decision",
        error="Pick a decision from the list.",
    )
    sheet.add_data_validation(validation)
    validation.add(f"B2:B{max(len(findings) + 1, 2)}")

    sheet.freeze_panes = "B2"


# ---------------------------------------------------------------- build


def _neutralise_formulas(workbook) -> None:
    """Store every '='-leading string as text, never as a formula.

    The workbook carries untrusted text: evidence values from source data and
    the executive brief, which is model output. openpyxl turns any string that
    starts with '=' into a live formula, so an evidence field or a prompt
    injection could plant =HYPERLINK(...) or =WEBSERVICE(...) in a reviewer's
    Excel. The workbook generates no formulas of its own (dropdowns live in data
    validation, not cells), so every formula cell is injected and safe to
    downgrade. Forcing the type keeps the text exact -- prefixing an apostrophe
    would corrupt the brief's markdown bullets.
    """
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"



def build(run_id: str | None = None, output_path: Path | None = None) -> Path:
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    output_path = output_path or EXCEL_DIR / "Finance_Review_Workbook.xlsx"

    with connect() as conn:
        if run_id is None:
            with conn.cursor() as cur:
                cur.execute("""
                    SELECT run_id FROM capstone.fact_pipeline_runs
                    WHERE status = 'SUCCESS' ORDER BY started_at DESC LIMIT 1
                """)
                row = cur.fetchone()
                if row is None:
                    raise SystemExit("no successful run found -- run the pipeline first")
                run_id = row[0]

        findings = fetch_open_findings(conn, run_id)
        breaks = fetch_reconciliation(conn, run_id)
        totals = fetch_control_totals(conn, run_id)

    fp_likelihood = load_fp_likelihood(BASE_DIR / "outputs")
    for finding in findings:
        finding["ai_fp_likelihood"] = fp_likelihood.get(finding["finding_id"])

    workbook = Workbook()
    workbook.remove(workbook.active)

    _sheet_instructions(workbook, run_id, len(findings))
    _sheet_review_queue(workbook, findings)
    _sheet_finding_evidence(workbook, findings)
    _sheet_decision_template(workbook, findings)
    _sheet_reconciliation(workbook, breaks)
    _sheet_control_totals(workbook, totals, run_id)
    _sheet_management_summary(workbook, run_id)
    _sheet_lists_and_rules(workbook)
    _neutralise_formulas(workbook)

    workbook.save(output_path)
    logger.info("workbook written to %s (%d open cases)", output_path, len(findings))
    return output_path


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    path = build()
    print(json.dumps({"workbook": str(path)}, indent=2))


if __name__ == "__main__":
    main()
