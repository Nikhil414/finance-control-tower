"""Verify the review importer refuses bad input and records good input honestly.

The importer is a gate, so most of these tests assert that something is
REJECTED. A control system's review step is only worth having if it cannot be
fed nonsense.

Requires a completed pipeline run. Run: pytest tests/test_import_reviews.py
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest

SRC = Path(__file__).resolve().parent.parent / "src"
sys.path.insert(0, str(SRC))

from db import connect  # noqa: E402
from import_reviews import ReviewFileRejected, import_file  # noqa: E402

COLUMNS = [
    "finding_id", "decision", "reviewer", "reviewed_at", "reviewer_notes",
    "approved_action", "confirmed_loss", "recovered_amount",
    "false_positive_reason", "resolution_date",
]


@pytest.fixture()
def conn():
    """A connection whose work is rolled back, so tests leave the case model
    exactly as they found it."""
    connection = connect()
    connection.autocommit = False
    yield connection
    connection.rollback()
    connection.close()


@pytest.fixture()
def run_id(conn):
    with conn.cursor() as cur:
        cur.execute("""
            SELECT run_id FROM capstone.fact_pipeline_runs
            WHERE status = 'SUCCESS' ORDER BY started_at DESC LIMIT 1
        """)
        row = cur.fetchone()
    if row is None:
        pytest.skip("no successful run -- run the pipeline first")
    return row[0]


@pytest.fixture()
def open_findings(conn, run_id):
    with conn.cursor() as cur:
        cur.execute("""
            SELECT finding_id, risk_amount FROM capstone.fact_findings
            WHERE last_seen_run_id = %s AND status = 'OPEN'
            ORDER BY priority_score DESC LIMIT 8
        """, (run_id,))
        found = cur.fetchall()
    if len(found) < 4:
        pytest.skip("not enough open findings to exercise the importer")
    return found


def write_csv(path: Path, rows: list[dict]) -> Path:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in COLUMNS})
    return path


def decision_row(finding_id: str, **overrides) -> dict:
    row = {
        "finding_id": finding_id,
        "decision": "VALID_EXCEPTION",
        "reviewer": "analyst.test",
        "reviewed_at": "2026-09-12",
        "reviewer_notes": "Verified against source evidence",
        "confirmed_loss": "1000.00",
    }
    row.update(overrides)
    return row


def expect_rejection(path, conn, run_id, fragment: str):
    with pytest.raises(ReviewFileRejected) as rejection:
        import_file(path, run_id=run_id, conn=conn)
    problems = " | ".join(rejection.value.problems)
    assert fragment in problems, f"expected '{fragment}' in: {problems}"


# ---------------------------------------------------------------- happy path

def test_valid_decisions_import(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0]),
        decision_row(open_findings[1][0], decision="MORE_INFORMATION_REQUIRED",
                     confirmed_loss=""),
    ])
    result = import_file(path, run_id=run_id, conn=conn)
    assert result["decisions_applied"] == 2


def test_blank_rows_are_skipped_not_rejected(tmp_path, conn, run_id, open_findings):
    """The template ships one row per open case; most stay empty."""
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0]),
        {"finding_id": open_findings[1][0]},
        {},
    ])
    result = import_file(path, run_id=run_id, conn=conn)
    assert result["decisions_applied"] == 1


def test_pickup_is_recorded_as_its_own_event(tmp_path, conn, run_id, open_findings):
    """A case must not appear to leap from untouched to decided."""
    finding_id = open_findings[0][0]
    import_file(write_csv(tmp_path / "d.csv", [decision_row(finding_id)]),
                run_id=run_id, conn=conn)

    with conn.cursor() as cur:
        cur.execute("""
            SELECT e.event_seq, e.from_status, e.to_status, e.actor_type
            FROM capstone.fact_finding_events e
            JOIN capstone.fact_findings f USING (finding_key)
            WHERE f.finding_id = %s ORDER BY e.event_seq
        """, (finding_id,))
        events = cur.fetchall()

    assert [event[1] for event in events] == [None, "OPEN", "UNDER_REVIEW"]
    assert [event[2] for event in events] == ["OPEN", "UNDER_REVIEW", "VALID_EXCEPTION"]
    assert events[0][3] == "SYSTEM"
    assert events[1][3] == "HUMAN" and events[2][3] == "HUMAN"


def test_review_is_recorded_against_the_finding(tmp_path, conn, run_id, open_findings):
    finding_id = open_findings[0][0]
    import_file(write_csv(tmp_path / "d.csv", [
        decision_row(finding_id, confirmed_loss="2500.00", recovered_amount="1000.00"),
    ]), run_id=run_id, conn=conn)

    with conn.cursor() as cur:
        cur.execute("""
            SELECT decision, reviewer, confirmed_loss, recovered_amount
            FROM capstone.vw_current_review cr
            JOIN capstone.fact_findings f USING (finding_key)
            WHERE f.finding_id = %s
        """, (finding_id,))
        decision, reviewer, loss, recovered = cur.fetchone()

    assert decision == "VALID_EXCEPTION"
    assert reviewer == "analyst.test"
    assert float(loss) == 2500.00
    assert float(recovered) == 1000.00


def test_re_review_supersedes_without_erasing(tmp_path, conn, run_id, open_findings):
    """Append-only means a correction adds a row. Both survive; the view shows
    the later one."""
    finding_id = open_findings[0][0]
    import_file(write_csv(tmp_path / "a.csv", [decision_row(finding_id)]),
                run_id=run_id, conn=conn)
    import_file(write_csv(tmp_path / "b.csv", [
        decision_row(finding_id, decision="MORE_INFORMATION_REQUIRED",
                     confirmed_loss="", reviewed_at="2026-09-13",
                     reviewer_notes="Reopened; gateway report contradicts the first read"),
    ]), run_id=run_id, conn=conn)

    with conn.cursor() as cur:
        cur.execute("""
            SELECT COUNT(*) FROM capstone.fact_reviews r
            JOIN capstone.fact_findings f USING (finding_key)
            WHERE f.finding_id = %s
        """, (finding_id,))
        assert cur.fetchone()[0] == 2

        cur.execute("""
            SELECT decision FROM capstone.vw_current_review cr
            JOIN capstone.fact_findings f USING (finding_key)
            WHERE f.finding_id = %s
        """, (finding_id,))
        assert cur.fetchone()[0] == "MORE_INFORMATION_REQUIRED"


# ---------------------------------------------------------------- rejections

def test_missing_required_column_is_rejected(tmp_path, conn, run_id):
    path = tmp_path / "d.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["finding_id", "decision"])
        writer.writerow(["LK-20260912-00001", "VALID_EXCEPTION"])
    expect_rejection(path, conn, run_id, "missing required column")


def test_unknown_decision_value_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], decision="looks fine to me"),
    ])
    expect_rejection(path, conn, run_id, "is not a valid decision")


def test_unknown_finding_id_is_rejected(tmp_path, conn, run_id):
    path = write_csv(tmp_path / "d.csv", [decision_row("LK-19990101-99999")])
    expect_rejection(path, conn, run_id, "is not a finding in run")


def test_anonymous_decision_is_rejected(tmp_path, conn, run_id, open_findings):
    """An unattributed decision is not auditable, which defeats the point."""
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], reviewer=""),
    ])
    expect_rejection(path, conn, run_id, "reviewer is required")


def test_false_positive_without_reason_is_rejected(tmp_path, conn, run_id, open_findings):
    """An unexplained dismissal is how control systems quietly rot."""
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], decision="FALSE_POSITIVE",
                     confirmed_loss="", reviewer_notes=""),
    ])
    expect_rejection(path, conn, run_id, "requires 'false_positive_reason'")


def test_false_positive_with_confirmed_loss_is_rejected(tmp_path, conn, run_id, open_findings):
    """If money was lost, it was not a false positive."""
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], decision="FALSE_POSITIVE",
                     false_positive_reason="Duplicate of another case",
                     confirmed_loss="5000.00"),
    ])
    expect_rejection(path, conn, run_id, "cannot carry a confirmed_loss")


def test_valid_exception_without_notes_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], reviewer_notes=""),
    ])
    expect_rejection(path, conn, run_id, "requires 'reviewer_notes'")


def test_recovery_exceeding_confirmed_loss_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], confirmed_loss="1000.00",
                     recovered_amount="5000.00"),
    ])
    expect_rejection(path, conn, run_id, "exceeds confirmed_loss")


def test_recovery_without_confirmed_loss_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], confirmed_loss="", recovered_amount="500.00"),
    ])
    expect_rejection(path, conn, run_id, "recovered_amount given with no confirmed_loss")


def test_negative_amount_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], confirmed_loss="-100.00"),
    ])
    expect_rejection(path, conn, run_id, "cannot be negative")


def test_non_numeric_amount_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], confirmed_loss="about five grand"),
    ])
    expect_rejection(path, conn, run_id, "is not a number")


def test_unparseable_date_is_rejected(tmp_path, conn, run_id, open_findings):
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], reviewed_at="last Tuesday"),
    ])
    expect_rejection(path, conn, run_id, "is not a recognised date")


def test_same_finding_twice_in_one_file_is_rejected(tmp_path, conn, run_id, open_findings):
    """Two decisions in one file have no defined order, so which one wins would
    be arbitrary."""
    finding_id = open_findings[0][0]
    path = write_csv(tmp_path / "d.csv", [
        decision_row(finding_id),
        decision_row(finding_id, decision="FALSE_POSITIVE", confirmed_loss="",
                     false_positive_reason="Changed my mind"),
    ])
    expect_rejection(path, conn, run_id, "already decided on row")


def test_approval_cannot_skip_the_valid_exception_step(tmp_path, conn, run_id, open_findings):
    """Approval is a real control gate. A single spreadsheet entry must not be
    able to both raise and approve its own remediation."""
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], decision="ACTION_APPROVED",
                     approved_action="Raise debit note", confirmed_loss=""),
    ])
    expect_rejection(path, conn, run_id, "approval is a separate control step")


def test_missing_file_is_rejected(tmp_path, conn, run_id):
    expect_rejection(tmp_path / "absent.csv", conn, run_id, "not found")


# ---------------------------------------------------------------- atomicity

def test_one_bad_row_prevents_the_whole_file(tmp_path, conn, run_id, open_findings):
    """All-or-nothing. A half-applied decision file leaves the case model in a
    state nobody intended and nobody can easily identify."""
    with conn.cursor() as cur:
        cur.execute("SELECT COUNT(*) FROM capstone.fact_reviews")
        before = cur.fetchone()[0]

    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0]),
        decision_row(open_findings[1][0]),
        decision_row(open_findings[2][0], decision="NONSENSE"),
    ])

    with pytest.raises(ReviewFileRejected):
        import_file(path, run_id=run_id, conn=conn)

    with conn.cursor() as cur:
        cur.execute("SELECT COUNT(*) FROM capstone.fact_reviews")
        assert cur.fetchone()[0] == before, "a rejected file must import nothing"


def test_all_problems_are_reported_at_once(tmp_path, conn, run_id, open_findings):
    """Reporting one problem per run would make fixing a spreadsheet an
    afternoon of guesswork."""
    path = write_csv(tmp_path / "d.csv", [
        decision_row(open_findings[0][0], decision="NONSENSE"),
        decision_row("LK-19990101-99999"),
        decision_row(open_findings[1][0], reviewer=""),
    ])
    with pytest.raises(ReviewFileRejected) as rejection:
        import_file(path, run_id=run_id, conn=conn)
    assert len(rejection.value.problems) >= 3


# ---------------------------------------------------------------- workbook

def test_workbook_has_every_required_sheet(tmp_path):
    """Eight sheets per the capstone specification."""
    from openpyxl import load_workbook

    from build_review_workbook import build

    path = build(output_path=tmp_path / "wb.xlsx")
    workbook = load_workbook(path)

    assert workbook.sheetnames == [
        "Instructions", "Review Queue", "Finding Evidence", "Reviewer Decisions",
        "Reconciliation", "Control Totals", "Management Summary", "Lists and Rules",
    ]


def test_workbook_decision_sheet_matches_importer_columns(tmp_path):
    """If the template's columns drift from what the importer reads, every
    reviewer's file is rejected for a reason that is not their fault."""
    from openpyxl import load_workbook

    from build_review_workbook import build

    workbook = load_workbook(build(output_path=tmp_path / "wb.xlsx"))
    header = [cell.value for cell in workbook["Reviewer Decisions"][1]]
    assert header == COLUMNS


def test_workbook_stores_injected_formulas_as_text(tmp_path):
    """Evidence and the AI-written brief are untrusted. A value like
    =HYPERLINK(...) must reach the reviewer as text, not as a live formula."""
    from openpyxl import Workbook, load_workbook

    from build_review_workbook import _neutralise_formulas

    payload = '=HYPERLINK("http://evil.example","click")'
    workbook = Workbook()
    workbook.active.append(["F-1", payload, "- markdown bullet", 42])
    _neutralise_formulas(workbook)
    workbook.save(tmp_path / "wb.xlsx")

    row = next(load_workbook(tmp_path / "wb.xlsx").active.iter_rows())
    assert row[1].data_type == "s"
    assert row[1].value == payload
    assert row[2].value == "- markdown bullet"
    assert row[3].value == 42


def test_workbook_control_totals_match_the_database(tmp_path, conn, run_id):
    """The workbook must be tieable back to Postgres, or it is an unverifiable
    extract."""
    from openpyxl import load_workbook

    from build_review_workbook import build

    workbook = load_workbook(build(output_path=tmp_path / "wb.xlsx"))
    totals = {
        row[0].value: row[1].value
        for row in workbook["Control Totals"].iter_rows(min_row=2)
        if row[0].value
    }

    with conn.cursor() as cur:
        cur.execute("SELECT COUNT(*) FROM capstone.fact_findings WHERE last_seen_run_id = %s",
                    (run_id,))
        expected = cur.fetchone()[0]

    assert totals["Findings in this run"] == expected
